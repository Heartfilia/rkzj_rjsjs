#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""docx 格式普查：判断每份答案详解属于哪种排版，能不能直接抽。

只读 data/_docx/*.docx，输出 data/docx-survey.md 和 data/docx-survey.json。
"""
import json
import os
import re
import sys
from collections import Counter

from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data", "_docx")
OUT_MD = os.path.join(ROOT, "data", "docx-survey.md")
OUT_JSON = os.path.join(ROOT, "data", "docx-survey.json")

ANSWER_TAG = re.compile(r"【答案】\s*([A-D])")
PLAIN_ANS = re.compile(r"^\s*(\d{1,2})(?:\s*[—\-~]\s*(\d{1,2}))?\s*([A-D])(?:\s+([A-D]))?(?=[\s，。、　]|$)")
INLINE_NUM = re.compile(r"[（(](\d{1,2})[）)]")
OPT_LINE = re.compile(r"^[（(]?\d{1,2}[）)]?\s*A[.．、]")


def load(path):
    doc = Document(path)
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    cells = []
    for t in doc.tables:
        for row in t.rows:
            cells.append(" | ".join(c.text.strip() for c in row.cells))
    return paras, cells


def survey_one(fname, paras, cells):
    text = "\n".join(paras)
    n_tag = len(ANSWER_TAG.findall(text))
    n_jiexi = text.count("【解析】")
    plain = [m for p in paras for m in [PLAIN_ANS.match(p)] if m]
    nums = set()
    for m in plain:
        nums.add(int(m.group(1)))
        if m.group(2):
            nums.add(int(m.group(2)))
    inline_nums = set(int(x) for x in INLINE_NUM.findall(text))
    opt_lines = sum(1 for p in paras if OPT_LINE.match(p))

    if n_tag >= 20:
        style = "tagged"          # 【答案】/【解析】 分段
    elif len(plain) >= 20:
        style = "compact"         # 1 D 解析…
    elif opt_lines >= 20:
        style = "paper_only"      # 只有卷子，没答案
    else:
        style = "unknown"

    return {
        "file": fname,
        "paras": len(paras),
        "tables": len(cells),
        "n_answer_tag": n_tag,
        "n_analysis_tag": n_jiexi,
        "n_plain": len(plain),
        "n_option_lines": opt_lines,
        "q_nums_inline": len(inline_nums),
        "q_nums_plain": len(nums),
        "style": style,
        "sample": paras[:2],
    }


def main():
    if not os.path.isdir(SRC):
        sys.exit(f"没有 {SRC}")
    rows = []
    for fname in sorted(os.listdir(SRC)):
        if not fname.endswith(".docx") or fname.startswith("~$"):
            continue
        try:
            paras, cells = load(os.path.join(SRC, fname))
        except Exception as e:                       # noqa: BLE001
            rows.append({"file": fname, "error": str(e)[:80]})
            continue
        rows.append(survey_one(fname, paras, cells))

    with open(OUT_JSON, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)

    lines = ["# 答案详解 docx 格式普查", "",
             "| 文件 | 段落 | 【答案】 | 【解析】 | 紧凑式 | 选项行 | 判定 |",
             "|---|---|---|---|---|---|---|"]
    for r in rows:
        if "error" in r:
            lines.append(f"| {r['file']} | - | - | - | - | - | ❌ {r['error']} |")
            continue
        lines.append(f"| {r['file'][:28]} | {r['paras']} | {r['n_answer_tag']} | "
                     f"{r['n_analysis_tag']} | {r['n_plain']} | {r['n_option_lines']} | {r['style']} |")

    lines += ["", "## 判断口径", "",
              "- **tagged**：`【答案】X` + `【解析】`，题干/选项/答案/解析都在同一份里，最好处理",
              "- **compact**：`1 D 解析…` 挤在一段，多空写成 `2—3B C`，通常只有答案没题干",
              "- **paper_only**：只有题目没有答案，说明拿错文件了", "",
              "## 每份的首段样本", ""]
    for r in rows:
        if r.get("sample"):
            lines.append(f"- **{r['file'][:30]}**：{r['sample'][0][:70]}")
    with open(OUT_MD, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    print("\n".join(lines[:12]))
    print(f"\n写出 {OUT_MD}")


if __name__ == "__main__":
    main()
